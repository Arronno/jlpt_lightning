import hashlib
import io
import logging
from collections import Counter
from zipfile import BadZipFile, ZipFile

from django.core.cache import cache
from django.db import transaction
from django.db.models import Count
from django.utils import timezone
from openpyxl import load_workbook

from .models import Dataset, Entry, ImportBatch, Lesson, normalize

FIELDS = ("serial", "lesson", "kana", "english", "example", "example_eng")
TEXT_FIELDS = ("kana", "english", "example", "example_eng")


def preview_import(raw, filename, level, key):
    errors, rows, notes = [], [], []
    if len(raw) > 10 * 1024 * 1024:
        raise ValueError("Workbooks must be smaller than 10 MB.")
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 50 MB.")
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Please select a valid .xlsx workbook.") from exc
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header = tuple(str(v or "").strip().lower() for v in next(iterator, ()))
        if header[:6] != FIELDS:
            errors.append("The first sheet must start with: " + ", ".join(FIELDS))
        else:
            seen = set()
            positions = Counter()
            for number, values in enumerate(iterator, 2):
                if number > 30001:
                    errors.append("A workbook may contain at most 30,000 entries.")
                    break
                if not any(v is not None for v in values):
                    continue
                cells = dict(zip(FIELDS, ("" if v is None else str(v) for v in values[:6])))
                try:
                    serial = int(cells["serial"])
                    if serial <= 0 or serial in seen:
                        raise ValueError
                    if any(not cells[field].strip() for field in FIELDS[:4]):
                        raise ValueError
                    if any(len(value) > 10000 or value.startswith("=") for value in cells.values()):
                        raise ValueError
                    if len(cells["lesson"]) > 100:
                        raise ValueError
                except (ValueError, KeyError):
                    errors.append(
                        f"Row {number}: required fields, positive unique serial, and literal text are required."
                    )
                    continue
                seen.add(serial)
                positions[cells["lesson"]] += 1
                cells.update(serial=serial, position=positions[cells["lesson"]])
                rows.append(cells)
        for sheet in workbook.worksheets[1:]:
            for row in sheet.iter_rows(values_only=True):
                notes.append(" | ".join(str(v) for v in row if v is not None))
                if len(notes) > 2000:
                    break
    finally:
        workbook.close()
    dataset = Dataset.objects.filter(key=key).first()
    if dataset and dataset.level != level:
        errors.append("This dataset key belongs to a different JLPT level.")
    if not rows:
        errors.append("The workbook contains no valid entries.")
    existing = (
        {e.serial: e for e in Entry.objects.filter(dataset=dataset).select_related("lesson")}
        if dataset
        else {}
    )
    additions, changes, unchanged = [], [], 0
    for row in rows:
        entry = existing.get(row["serial"])
        if not entry:
            additions.append(row["serial"])
        elif (
            any(getattr(entry, f) != row[f] for f in TEXT_FIELDS)
            or entry.lesson.label != row["lesson"]
        ):
            changes.append(
                {
                    "serial": row["serial"],
                    "before": {
                        **{f: getattr(entry, f) for f in TEXT_FIELDS},
                        "lesson": entry.lesson.label,
                    },
                    "after": row,
                }
            )
        else:
            unchanged += 1
    missing = sorted(set(existing) - {row["serial"] for row in rows})
    return ImportBatch.objects.create(
        dataset_key=key,
        level=level,
        filename=filename[:255],
        digest=hashlib.sha256(raw).hexdigest(),
        errors=errors,
        base_revision=dataset.revision if dataset else 0,
        payload={
            "rows": rows,
            "notes": "\n".join(notes),
            "added": additions,
            "changed": changes,
            "unchanged": unchanged,
            "missing": missing,
        },
    )


def apply_import(batch_id, backup=True):
    from .storage import backup_database

    if backup:
        backup_database("before-import")
    with transaction.atomic():
        batch = ImportBatch.objects.get(pk=batch_id)
        if batch.applied_at:
            return batch
        if batch.errors:
            raise ValueError("Fix workbook errors and upload it again before importing.")
        dataset, _ = Dataset.objects.get_or_create(
            key=batch.dataset_key, defaults={"level": batch.level}
        )
        if dataset.revision != batch.base_revision:
            raise ValueError(
                "This preview is out of date. Upload the workbook again for a fresh preview."
            )
        labels = list(dict.fromkeys(row["lesson"] for row in batch.payload["rows"]))
        lessons = {}
        for i, label in enumerate(labels):
            lessons[label], _ = Lesson.objects.get_or_create(
                dataset=dataset, label=label, defaults={"position": i}
            )
        existing = {e.serial: e for e in Entry.objects.filter(dataset=dataset)}
        positions = Counter()
        for entry in existing.values():
            positions[entry.lesson_id] = max(positions[entry.lesson_id], entry.position)
        added, changed = [], []
        for row in batch.payload["rows"]:
            entry = existing.get(row["serial"])
            new = entry is None
            if new:
                entry = Entry(dataset=dataset, serial=row["serial"])
            for field in TEXT_FIELDS:
                setattr(entry, field, row[field])
            target_lesson = lessons[row["lesson"]]
            if new or entry.lesson_id != target_lesson.pk:
                positions[target_lesson.pk] += 1
                entry.position = positions[target_lesson.pk]
            entry.lesson = target_lesson
            entry.batch = (entry.position - 1) // 20 + 1
            entry.search = normalize(row["kana"] + " " + row["english"]).casefold()
            (added if new else changed).append(entry)
        Entry.objects.bulk_create(added)
        Entry.objects.bulk_update(changed, [*TEXT_FIELDS, "lesson", "position", "batch", "search"])
        dataset.notes = batch.payload["notes"]
        dataset.revision += 1
        dataset.save()
        batch.applied_at = timezone.now()
        batch.save(update_fields=["applied_at"])
    logging.getLogger(__name__).info(
        "Imported dataset %s: %d additions", batch.dataset_key, len(added)
    )
    return batch


def filter_entries(filters):
    entries = Entry.objects.select_related("dataset", "lesson")
    for key, lookup in [
        ("level", "dataset__level"),
        ("lesson", "lesson_id"),
        ("batch", "batch"),
        ("collection", "collections__id"),
    ]:
        if filters.get(key):
            entries = entries.filter(**{lookup: filters[key]})
    if filters.get("q"):
        entries = entries.filter(search__contains=normalize(filters["q"]).casefold())
    if filters.get("bookmarked"):
        entries = entries.filter(annotations__profile_id=1, annotations__bookmarked=True)
    if filters.get("tag"):
        entries = entries.filter(
            annotations__profile_id=1, annotations__tags__icontains=filters["tag"]
        )
    return entries.distinct()


def catalog_summary():
    revision = list(Dataset.objects.order_by("pk").values_list("pk", "revision"))
    key = "catalog:" + hashlib.sha256(repr(revision).encode()).hexdigest()
    result = cache.get(key)
    if result is None:
        result = [
            {
                "level": level,
                "count": Entry.objects.filter(dataset__level=level).count(),
                "lessons": list(
                    Lesson.objects.filter(dataset__level=level)
                    .annotate(count=Count("entry"))
                    .filter(count__gt=0)
                    .values("id", "label", "count")
                ),
            }
            for level in range(5, 0, -1)
        ]
        cache.set(key, result)
    return result
